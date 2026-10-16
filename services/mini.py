"""DrDemic Mini: validation, calculation, ranking, history diff, xlsx. No Flask."""
import io, re, secrets

from core.security import hash_password_fast
from services.access import normalize_phone
from services.students_import import _parse_with_headers

TYPES = ("marks", "primary", "secondary", "advanced")
DEFAULT_GRADES = [{"min": 80, "grade": "A"}, {"min": 70, "grade": "B"}, {"min": 60, "grade": "C"},
                  {"min": 50, "grade": "D"}, {"min": 0, "grade": "F"}]
O_LEVEL = [(80, "A", 1), (70, "B", 2), (60, "C", 3), (50, "D", 4), (0, "F", 5)]
A_LEVEL = [(80, "A", 1), (70, "B", 2), (60, "C", 3), (50, "D", 4), (40, "E", 5), (35, "S", 6), (0, "F", 7)]
O_DIV = [(7, 17, "I"), (18, 21, "II"), (22, 25, "III"), (26, 33, "IV"), (34, 50, "0")]
A_DIV = [(3, 9, "I"), (10, 12, "II"), (13, 17, "III"), (18, 19, "IV"), (20, 21, "0")]
_SKIP = {"sn", "no", "total", "average", "avg", "position", "pos", "grade", "grd", "points", "division"}


def new_id(prefix="x"):
    return prefix + secrets.token_hex(4)


def default_data():
    names = ("English", "Kiswahili", "History", "Geography", "Subject 5", "Subject 6")
    return {"subjects": [{"id": new_id("s"), "name": n} for n in names], "rows": []}


def clean_data(d):
    """-> (clean_data, None) or (None, error message)."""
    if not isinstance(d, dict):
        return None, "Invalid sheet data"
    subs, seen = [], set()
    for s in (d.get("subjects") or [])[:40]:
        sid = str(s.get("id") or new_id("s"))[:20]
        if sid in seen:
            continue
        seen.add(sid)
        subs.append({"id": sid, "name": str(s.get("name") or "").strip()[:40]})
    rows, rseen = [], set()
    for r in (d.get("rows") or [])[:2000]:
        name = str(r.get("name") or "").strip()[:120]
        phone = str(r.get("phone") or "").strip()[:30]
        marks = {}
        for sid, v in (r.get("marks") or {}).items():
            if sid not in seen or v is None or str(v).strip() == "":
                continue
            try:
                x = float(v)
            except (TypeError, ValueError):
                return None, f"Invalid mark for {name or 'a student'}"
            if not (0 <= x <= 100):
                return None, f"Marks must be between 0 and 100 ({name or 'a student'})"
            marks[sid] = round(x, 2)
        if not name and not phone and not marks:
            continue
        rid = str(r.get("id") or new_id("r"))[:20]
        if rid in rseen:
            rid = new_id("r")
        rseen.add(rid)
        rows.append({"id": rid, "name": name, "phone": phone, "marks": marks})
    return {"subjects": subs, "rows": rows}, None


def clean_grades(rules):
    out = []
    try:
        for g in rules:
            m, name = float(g["min"]), str(g["grade"]).strip()
            if not (0 <= m <= 100) or not (1 <= len(name) <= 3):
                return None
            out.append({"min": m, "grade": name})
    except (TypeError, ValueError, KeyError):
        return None
    if not out or len(out) > 12:
        return None
    return sorted(out, key=lambda g: -g["min"])


def grade_for(rules, score):
    if score is None:
        return "-"
    for g in rules:
        if score >= g["min"]:
            return g["grade"]
    return rules[-1]["grade"]


def _points(table, score):
    for m, _g, p in table:
        if score >= m:
            return p
    return table[-1][2]


def compute(data, sheet_type, rules):
    """Rows with total/average/grade/points/division/position, sorted best first (ties share a position)."""
    ids = [s["id"] for s in data["subjects"] if s["name"]]
    out = []
    for r in data["rows"]:
        vals = [r["marks"][i] for i in ids if i in r["marks"]]
        total = round(sum(vals), 2) if vals else None
        avg = round(total / len(vals), 2) if vals else None
        row = dict(r, total=total, average=avg, grade=grade_for(rules, avg), points=None, division=None)
        if sheet_type in ("secondary", "advanced"):
            table, divs, need = (O_LEVEL, O_DIV, 7) if sheet_type == "secondary" else (A_LEVEL, A_DIV, 3)
            pts = sorted(_points(table, v) for v in vals)
            if len(pts) < need:
                row["points"] = row["division"] = "INC"
            else:
                row["points"] = sum(pts[:need])
                row["division"] = next((d for lo, hi, d in divs if lo <= row["points"] <= hi), "0")
        out.append(row)
    ranked = sorted(out, key=lambda r: (r["average"] is None, -(r["average"] or 0), r["name"].lower()))
    prev = None
    for i, r in enumerate(ranked):
        if r["average"] is None:
            r["position"] = None
            continue
        r["position"] = prev["position"] if prev and prev["average"] == r["average"] else i + 1
        prev = r
    return ranked


def diff(old, new, old_meta, new_meta):
    ch = []
    osub = {s["id"]: s["name"] for s in old["subjects"]}
    nsub = {s["id"]: s["name"] for s in new["subjects"]}
    for sid, n in nsub.items():
        if sid not in osub:
            ch.append({"what": "subject_added", "new": n})
        elif osub[sid] != n:
            ch.append({"what": "subject_renamed", "old": osub[sid], "new": n})
    for sid, n in osub.items():
        if sid not in nsub:
            ch.append({"what": "subject_removed", "old": n})
    orows = {r["id"]: r for r in old["rows"]}
    nrows = {r["id"]: r for r in new["rows"]}
    for rid, r in nrows.items():
        o = orows.get(rid)
        if o is None:
            ch.append({"what": "student_added", "student": r["name"],
                       "marks": {nsub.get(k, k): v for k, v in r["marks"].items()}})
            continue
        for f in ("name", "phone"):
            if o[f] != r[f]:
                ch.append({"what": f, "student": r["name"] or o["name"], "old": o[f], "new": r[f]})
        for sid in nsub:
            a, b = o["marks"].get(sid), r["marks"].get(sid)
            if a != b:
                ch.append({"what": "mark", "student": r["name"], "subject": nsub[sid], "old": a, "new": b})
    for rid, o in orows.items():
        if rid not in nrows:
            ch.append({"what": "student_removed", "student": o["name"]})
    for k in ("school_name", "exam_name", "sheet_type", "grade_rules"):
        if old_meta[k] != new_meta[k]:
            ch.append({"what": k, "old": old_meta[k], "new": new_meta[k]})
    return ch[:1000]


def sync_recipients(cur, sheet_id, rows, create_accounts):
    """Rebuild which parent phones can see this sheet; optionally create their accounts.
    First password = the phone in local format (0712345678), same convention as the main app."""
    phones = {normalize_phone(r["phone"]) for r in rows} - {None}
    cur.execute("DELETE FROM mini_recipients WHERE sheet_id=%s", (sheet_id,))
    for p in phones:
        cur.execute("INSERT INTO mini_recipients(sheet_id,phone) VALUES(%s,%s)", (sheet_id, p))
        if create_accounts:
            cur.execute("""INSERT INTO mini_users(username,password,role,must_change_password)
                           VALUES(%s,%s,'parent',1) ON CONFLICT DO NOTHING""",
                        (p, hash_password_fast("0" + p[3:])))
    return len(phones)


def parse_import(file_obj, filename):
    headers, rows = _parse_with_headers(file_obj, filename)
    if not headers or not rows:
        raise ValueError("File is empty or has no data rows")
    c = {h: re.sub(r"[^a-z0-9]", "", h.lower()) for h in headers}
    name_h = next((h for h in headers if any(k in c[h] for k in ("name", "student", "jina"))), headers[0])
    phone_h = next((h for h in headers if h != name_h and
                    any(k in c[h] for k in ("phone", "account", "simu", "contact", "mobile"))), None)
    subs = [h for h in headers if h not in (name_h, phone_h) and c[h] not in _SKIP]
    subjects = [{"id": new_id("s"), "name": h[:40]} for h in subs]
    out = []
    for r in rows:
        marks = {}
        for s, h in zip(subjects, subs):
            try:
                x = float(r.get(h, "")) if r.get(h, "") != "" else None
            except ValueError:
                x = None
            if x is not None and 0 <= x <= 100:
                marks[s["id"]] = x
        out.append({"id": new_id("r"), "name": r.get(name_h, ""),
                    "phone": r.get(phone_h, "") if phone_h else "", "marks": marks})
    return {"subjects": subjects, "rows": out}


def _style(ws, row_no):
    from openpyxl.styles import Font, PatternFill, Alignment
    for c in ws[row_no]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A6FA8")
        c.alignment = Alignment(horizontal="center")


def sheet_xlsx(meta, subjects, ranked):
    """Same shape as the existing score sheet: no phone column."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Score Sheet"
    ws.append([meta["school_name"] or "School"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([meta["exam_name"] or "Examination"]); ws.append([])
    extra = ["Points", "Division"] if meta["sheet_type"] in ("secondary", "advanced") else []
    ws.append(["S/N", "Student"] + [s["name"] for s in subjects] + ["Total", "Average", "Position", "Grade"] + extra)
    _style(ws, 4)
    for i, r in enumerate(ranked, 1):
        ws.append([i, r["name"]] + [r["marks"].get(s["id"]) for s in subjects] +
                  [r["total"], r["average"], r["position"], r["grade"]] +
                  ([r["points"], r["division"]] if extra else []))
    ws.column_dimensions["B"].width = 28
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def template_xlsx(subjects):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "DrDemic Sheet"
    ws.append(["S/N", "Student Name", "Account/Phone Number"] + [s["name"] for s in subjects if s["name"]])
    _style(ws, 1)
    ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 24
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf